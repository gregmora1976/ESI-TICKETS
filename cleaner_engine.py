from __future__ import annotations

from io import BytesIO
import hashlib
import json
import re
import unicodedata
from typing import Dict, List, Tuple, Iterable


def _text(value) -> str:
    if value is None:
        return ""
    value = str(value).replace("\u00a0", " ")
    value = value.replace("\r", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\s*\n\s*", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def _norm(value) -> str:
    s = _text(value).casefold()
    s = "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _slug(value) -> str:
    return re.sub(r"[^a-z0-9]+", "-", _norm(value)).strip("-") or "profil"


def _profile_key(mode: str, name: str) -> str:
    return f"{_slug(mode)}:{_slug(name)}"


def _owner_from_text(text: str, aliases: Iterable[str] | None = None) -> str:
    lines = [x.strip() for x in (text or "").replace("\r", "").splitlines() if x.strip()]
    labels = list(aliases or []) + ["consignor/shipper", "consignor", "shipper", "owner", "lender"]
    labels = list(dict.fromkeys(_norm(x) for x in labels if _norm(x)))
    for i, line in enumerate(lines):
        n = _norm(line)
        for nl in labels:
            if n.startswith(nl):
                rest = re.sub(r"^[^:]{0,60}:\s*", "", line).strip()
                if rest and _norm(rest) != nl:
                    return _text(rest)
                if i + 1 < len(lines):
                    return _text(lines[i + 1])
    m = re.search(r"property of\s+(.+?)(?:\.|\n| and are )", text or "", re.I | re.S)
    if m:
        return _text(m.group(1))
    return ""


def _money(value) -> Tuple[str, float | None]:
    raw = _text(value)
    if not raw:
        return "", None
    cleaned = re.sub(r"[^0-9,.-]", "", raw)
    if not cleaned:
        return raw, None
    # Documents de référence : virgule = séparateur de milliers, point = décimales.
    cleaned = cleaned.replace(",", "")
    try:
        number = float(cleaned)
    except Exception:
        return raw, None
    if number.is_integer():
        return f"{int(number):,}", number
    return f"{number:,.2f}", number


def _dimensions(value, default_unit: str = "") -> str:
    s = _text(value)
    if not s:
        return ""
    s = re.sub(r"^overall\s*:\s*", "", s, flags=re.I)
    s = s.replace('"', '″').replace("”", "″").replace("“", "″")
    s = re.sub(r"\s+[xX]\s+", " × ", s)
    s = re.sub(r"(?<=\d)[xX](?=\d)", " × ", s)
    s = re.sub(r"\s+", " ", s).strip()
    if default_unit and not re.search(r"\b(?:cm|mm|m|in|inch|inches)\b|[″']", s, re.I):
        s += f" {default_unit}"
    return s


def _license_from_note(note) -> str:
    s = _text(note)
    if not s:
        return ""
    s = re.sub(r"(?<=-)\s+(?=[A-Za-z0-9])", "", s)
    patterns = [
        r"\b\d{2}CA\d{5}/[A-Z0-9]+\b",
        r"\b\d{4}-\d{2}-\d{2}-\d{2}-\d{3}\b",
        r"\b[A-Z]{1,5}-?\d{4,}(?:/[A-Z0-9-]+)?\b",
    ]
    for pattern in patterns:
        m = re.search(pattern, s, re.I)
        if m:
            return m.group(0)
    if re.search(r"\b(?:permit|licen[cs]e|cites)\b", s, re.I):
        return s
    return ""


def _base_row(owner: str, source_file: str, source_page: int | str, reference: str = "") -> Dict:
    return {
        "owner": owner or "",
        "title": "",
        "medium": "",
        "hs_code": "",
        "license": "",
        "dimensions": "",
        "value_usd": "",  # Jamais rempli depuis CAD sans conversion utilisateur.
        "country_origin": "",
        "reference": reference or "",
        "source_file": source_file,
        "source_page": source_page,
        "source_value": "",
        "source_value_number": None,
        "source_currency": "",
        "source_headers": {},
    }


def _profile_name(full_text: str) -> str:
    n = _norm(full_text)
    if "museum of anthropology at ubc" in n and "object id" in n:
        return "MOA - Pro Forma"
    if "herve curat" in n and "item id" in n:
        return "Hervé Curat - Pro Forma"
    if "museum of vancouver" in n:
        return "Museum of Vancouver - Pro Forma"
    if "meghann o brien" in n and "spirit of shape" in n:
        return "Meghann O'Brien - Pro Forma"
    return "Générique"


_HEADER_ALIASES = {
    "owner": ["owner", "lender", "consignor", "consignor shipper", "shipper", "proprietaire", "preteur"],
    "reference": ["object id", "catalogue number", "catalog number", "item id", "reference", "ref", "inventory number", "accession number"],
    "title": ["object name title", "object name", "title", "description", "artwork", "artwork title"],
    "medium": ["medium", "materials", "material", "technique"],
    "hs_code": ["hs code", "tariff code", "commodity code", "customs code"],
    "license": ["license", "licence", "permit", "export permit", "cites"],
    "dimensions": ["dimensions", "dims", "size", "dimensions cm"],
    "country_origin": ["country of origin", "place of origin", "origin", "pays d origine"],
    "value_source": ["value for insurance cad", "value for insurance", "value cad", "value usd", "declared value", "customs value", "value"],
}


def _aliases_from_rules(learned_rules: List[Dict] | None, profile_key: str = "") -> Dict[str, List[str]]:
    aliases = {k: list(v) for k, v in _HEADER_ALIASES.items()}
    for rule in learned_rules or []:
        if not isinstance(rule, dict) or rule.get("active") is False:
            continue
        rule_profile = _text(rule.get("profile_key"))
        if rule_profile and profile_key and rule_profile != profile_key:
            continue
        if rule_profile and not profile_key:
            continue
        target = _text(rule.get("target_field"))
        alias = _text(rule.get("source_alias"))
        if target and alias:
            aliases.setdefault(target, [])
            if _norm(alias) not in {_norm(x) for x in aliases[target]}:
                aliases[target].append(alias)
    return aliases


def _match_header(value: str, aliases_map: Dict[str, List[str]] | None = None) -> str:
    n = _norm(value)
    if not n:
        return ""
    for field, aliases in (aliases_map or _HEADER_ALIASES).items():
        for alias in aliases:
            a = _norm(alias)
            if n == a or (len(a) >= 6 and a in n):
                return field
    return ""


def _currency_from_header(header: str, full_text: str) -> str:
    n = _norm(header)
    if "usd" in n:
        return "USD"
    if "cad" in n:
        return "CAD"
    t = _norm(full_text)
    if "cad" in t and "usd" not in t:
        return "CAD"
    if "usd" in t and "cad" not in t:
        return "USD"
    return ""


def _signature_terms_from_text(full_text: str) -> List[str]:
    terms = []
    for raw in (full_text or "").replace("\r", "").splitlines()[:120]:
        txt = _text(raw)
        n = _norm(txt)
        if not n or len(n) > 90:
            continue
        if any(k in n for k in (
            "pro forma", "invoice", "consignor", "shipper", "owner", "lender",
            "object id", "item id", "object name", "title", "dimensions", "value",
            "country of origin", "place of origin", "license", "permit", "cites",
        )):
            terms.append(n)
    return terms


def _signature_terms_pdf(pdf, full_text: str) -> List[str]:
    terms = _signature_terms_from_text(full_text)
    for page in list(pdf.pages)[:4]:
        try:
            tables = page.extract_tables() or []
        except Exception:
            tables = []
        for table in tables[:4]:
            for raw in (table or [])[:5]:
                for cell in raw or []:
                    txt = _text(cell)
                    n = _norm(txt)
                    if n and 2 <= len(n) <= 80 and not re.fullmatch(r"[0-9 .,/:-]+", n):
                        terms.append(n)
    return sorted(set(terms))[:120]


def _signature_terms_xlsx(wb) -> List[str]:
    terms = []
    for ws in wb.worksheets[:4]:
        terms.append(_norm(ws.title))
        for row_no in range(1, min(ws.max_row, 15) + 1):
            for col_no in range(1, min(ws.max_column, 40) + 1):
                txt = _text(ws.cell(row_no, col_no).value)
                n = _norm(txt)
                if n and 2 <= len(n) <= 80 and not re.fullmatch(r"[0-9 .,/:-]+", n):
                    terms.append(n)
    return sorted(set(x for x in terms if x))[:120]


def _signature_hash(signature: List[str]) -> str:
    raw = "|".join(sorted(set(_norm(x) for x in signature if _norm(x))))
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]


def _decode_signature(value) -> List[str]:
    if isinstance(value, list):
        return [_norm(x) for x in value if _norm(x)]
    if isinstance(value, str) and value.strip():
        try:
            parsed = json.loads(value)
            if isinstance(parsed, list):
                return [_norm(x) for x in parsed if _norm(x)]
        except Exception:
            return [_norm(x) for x in value.split("|") if _norm(x)]
    return []


def _signature_similarity(a: List[str], b: List[str]) -> float:
    sa = set(_norm(x) for x in a if _norm(x))
    sb = set(_norm(x) for x in b if _norm(x))
    if not sa or not sb:
        return 0.0
    # Jaccard enrichi par recouvrement du plus petit profil : robuste aux colonnes ajoutées.
    inter = len(sa & sb)
    jaccard = inter / len(sa | sb)
    coverage = inter / min(len(sa), len(sb))
    return round((jaccard * 0.45) + (coverage * 0.55), 4)


def _best_learned_profile(signature: List[str], learned_profiles: List[Dict] | None, threshold: float = 0.52):
    best = None
    best_score = 0.0
    for profile in learned_profiles or []:
        if not isinstance(profile, dict) or profile.get("active") is False:
            continue
        score = _signature_similarity(signature, _decode_signature(profile.get("signature")))
        if score > best_score:
            best = profile
            best_score = score
    if best is not None and best_score >= threshold:
        return best, best_score
    return None, best_score


def _extract_moa(pdf, owner: str, filename: str) -> List[Dict]:
    rows = []
    seen = set()
    for page_no, page in enumerate(pdf.pages, 1):
        for table in page.extract_tables() or []:
            if not table or max((len(r or []) for r in table), default=0) < 10:
                continue
            for raw in table:
                raw = list(raw or []) + [""] * 11
                ref = _text(raw[2])
                if not ref or _norm(ref) in {"object id", "object id #"}:
                    continue
                if not re.search(r"[A-Za-z0-9]", ref):
                    continue
                title = _text(raw[3])
                value = _text(raw[7])
                if not title and not value:
                    continue
                key = ref.casefold()
                if key in seen:
                    continue
                seen.add(key)
                display_value, number = _money(value)
                row = _base_row(owner, filename, page_no, ref)
                row.update({
                    "title": title,
                    "country_origin": _text(raw[5]),
                    "dimensions": _dimensions(raw[8]),
                    "license": _license_from_note(raw[10]),
                    "source_value": display_value,
                    "source_value_number": number,
                    "source_currency": "CAD",
                    "source_headers": {
                        "reference": "Object ID #", "title": "Object Name/Title",
                        "country_origin": "Place of Origin", "dimensions": "Dimensions",
                        "license": "Note", "value_source": "Value for Insurance CAD",
                    },
                })
                rows.append(row)
    return rows


def _extract_herve(pdf, owner: str, filename: str) -> List[Dict]:
    rows = []
    seen = set()
    for page_no, page in enumerate(pdf.pages, 1):
        for table in page.extract_tables() or []:
            if not table or max((len(r or []) for r in table), default=0) < 9:
                continue
            for raw in table:
                raw = list(raw or []) + [""] * 9
                seq = _text(raw[0])
                ref = _text(raw[2])
                if not seq.isdigit() or not ref or ref.casefold() in seen:
                    continue
                seen.add(ref.casefold())
                display_value, number = _money(raw[8])
                row = _base_row(owner, filename, page_no, ref)
                row.update({
                    "title": _text(raw[5]),
                    "dimensions": _dimensions(raw[3], "cm"),
                    "country_origin": _text(raw[4]),
                    "source_value": display_value,
                    "source_value_number": number,
                    "source_currency": "CAD",
                    "source_headers": {
                        "reference": "Item ID", "dimensions": "Dimensions (cm)",
                        "country_origin": "Place of Origin", "title": "Description",
                        "value_source": "Value (CAD)",
                    },
                })
                rows.append(row)
    return rows


def _extract_mov(pdf, owner: str, filename: str) -> List[Dict]:
    rows = []
    seen = set()

    def add(ref, title, value, origin, page_no):
        ref = _text(ref)
        title = _text(title)
        if not ref or not title or ref.casefold() in seen or not re.match(r"^AA\s+", ref, re.I):
            return
        seen.add(ref.casefold())
        display_value, number = _money(value)
        row = _base_row(owner, filename, page_no, ref)
        row.update({
            "title": title,
            "country_origin": _text(origin),
            "source_value": display_value,
            "source_value_number": number,
            "source_currency": "CAD",
            "source_headers": {
                "reference": "Catalogue Number", "title": "Object Name",
                "country_origin": "Place of Origin", "value_source": "Value for Insurance (CAD)",
            },
        })
        rows.append(row)

    for page_no, page in enumerate(pdf.pages, 1):
        for table in page.extract_tables() or []:
            for raw in table or []:
                raw = list(raw or []) + [""] * 5
                add(raw[1], raw[2], raw[3], raw[4], page_no)
        try:
            table = page.extract_table({
                "vertical_strategy": "text", "horizontal_strategy": "text",
                "intersection_tolerance": 5, "snap_tolerance": 3, "join_tolerance": 3,
                "edge_min_length": 3, "min_words_vertical": 1, "min_words_horizontal": 1,
            })
        except Exception:
            table = None
        if table:
            pending_title = []
            pending_origin = []
            for raw in table:
                raw = list(raw or []) + [""] * 9
                ref = _text(raw[5])
                title_piece = _text(raw[6])
                value = _text(raw[7])
                origin_piece = _text(raw[8])
                if ref and re.match(r"^AA\s+", ref, re.I):
                    title = " ".join(pending_title + ([title_piece] if title_piece else []))
                    origin = " ".join(pending_origin + ([origin_piece] if origin_piece else []))
                    add(ref, title, value, origin, page_no)
                    pending_title = []
                    pending_origin = []
                else:
                    if title_piece and "total value" not in _norm(title_piece):
                        pending_title.append(title_piece)
                    if origin_piece:
                        pending_origin.append(origin_piece)
    return rows


def _extract_meghann(pdf, owner: str, filename: str) -> List[Dict]:
    rows = []
    for page_no, page in enumerate(pdf.pages, 1):
        for table in page.extract_tables() or []:
            for raw in table or []:
                raw = list(raw or []) + [""] * 5
                dims = _text(raw[1])
                title = _text(raw[2])
                value = _text(raw[3])
                origin = _text(raw[4])
                if not title or _norm(title) in {"object name", "title"}:
                    continue
                if not value or not re.search(r"\d", value):
                    continue
                display_value, number = _money(value)
                row = _base_row(owner, filename, page_no)
                row.update({
                    "title": title,
                    "dimensions": _dimensions(dims),
                    "country_origin": origin,
                    "source_value": display_value,
                    "source_value_number": number,
                    "source_currency": "CAD",
                    "source_headers": {
                        "dimensions": "Dims", "title": "Object Name",
                        "value_source": "Value for Insurance (CAD)", "country_origin": "Place of Origin",
                    },
                })
                rows.append(row)
    return rows


def _apply_generic_fields(row: Dict, mapping: Dict[str, int], raw: List, header_values: Dict[str, str], currency: str):
    def cell(field):
        idx = mapping.get(field)
        return raw[idx] if idx is not None and idx < len(raw) else ""

    ref = _text(cell("reference"))
    row["reference"] = ref
    row["title"] = _text(cell("title"))
    row["medium"] = _text(cell("medium"))
    row["hs_code"] = _text(cell("hs_code"))
    row["license"] = _license_from_note(cell("license")) or _text(cell("license"))
    row["dimensions"] = _dimensions(cell("dimensions"))
    row["country_origin"] = _text(cell("country_origin"))
    display_value, number = _money(cell("value_source"))
    row["source_value"] = display_value
    row["source_value_number"] = number
    row["source_currency"] = currency
    row["value_usd"] = display_value if currency == "USD" else ""
    row["source_headers"] = dict(header_values)

    reserved = {"owner", "reference", "title", "medium", "hs_code", "license", "dimensions", "country_origin", "value_source"}
    for field in mapping:
        if field not in reserved:
            row[field] = _text(cell(field))
    return row


def _extract_generic_pdf(pdf, owner: str, filename: str, full_text: str, aliases_map: Dict[str, List[str]]) -> List[Dict]:
    rows = []
    seen = set()
    for page_no, page in enumerate(pdf.pages, 1):
        for table in page.extract_tables() or []:
            if not table:
                continue
            header_idx = None
            mapping = {}
            header_values = {}
            for i, raw in enumerate(table[:6]):
                candidate = {}
                candidate_values = {}
                for col, cell_value in enumerate(raw or []):
                    field = _match_header(cell_value, aliases_map)
                    if field and field not in candidate:
                        candidate[field] = col
                        candidate_values[field] = _text(cell_value)
                if "title" in candidate and len(candidate) >= 2:
                    header_idx = i
                    mapping = candidate
                    header_values = candidate_values
                    break
            if header_idx is None:
                continue
            currency = _currency_from_header(header_values.get("value_source", ""), full_text)
            for raw in table[header_idx + 1:]:
                raw = list(raw or [])
                title_idx = mapping.get("title")
                title = _text(raw[title_idx]) if title_idx is not None and title_idx < len(raw) else ""
                if not title:
                    continue
                ref_idx = mapping.get("reference")
                ref = _text(raw[ref_idx]) if ref_idx is not None and ref_idx < len(raw) else ""
                signature = (ref.casefold(), title.casefold(), page_no)
                if signature in seen:
                    continue
                seen.add(signature)
                row = _base_row(owner, filename, page_no, ref)
                rows.append(_apply_generic_fields(row, mapping, raw, header_values, currency))
    return rows


def extract_customs_pdf(filename: str, content: bytes, learned_profiles: List[Dict] | None = None, learned_rules: List[Dict] | None = None) -> Dict:
    import pdfplumber

    with pdfplumber.open(BytesIO(content)) as pdf:
        page_texts = [p.extract_text() or "" for p in pdf.pages]
        full_text = "\n".join(page_texts)
        signature = _signature_terms_pdf(pdf, full_text)
        builtin_profile = _profile_name(full_text)

        matched_profile = None
        learned_score = 0.0
        if builtin_profile == "Générique":
            matched_profile, learned_score = _best_learned_profile(signature, learned_profiles)

        if builtin_profile != "Générique":
            profile = builtin_profile
            profile_key = _profile_key("douanes", profile)
            profile_confidence = 1.0
            profile_source = "built-in"
        elif matched_profile:
            profile = _text(matched_profile.get("name")) or "Profil appris"
            profile_key = _text(matched_profile.get("profile_key")) or _profile_key("douanes", profile)
            profile_confidence = learned_score
            profile_source = "learned"
        else:
            profile = "Générique"
            profile_key = f"douanes:generic:{_signature_hash(signature)}"
            profile_confidence = learned_score
            profile_source = "generic"

        aliases_map = _aliases_from_rules(learned_rules, profile_key)
        owner = _owner_from_text(full_text, aliases_map.get("owner"))

        if builtin_profile == "MOA - Pro Forma":
            rows = _extract_moa(pdf, owner, filename)
        elif builtin_profile == "Hervé Curat - Pro Forma":
            rows = _extract_herve(pdf, owner, filename)
        elif builtin_profile == "Museum of Vancouver - Pro Forma":
            rows = _extract_mov(pdf, owner, filename)
        elif builtin_profile == "Meghann O'Brien - Pro Forma":
            rows = _extract_meghann(pdf, owner, filename)
        else:
            rows = _extract_generic_pdf(pdf, owner, filename, full_text, aliases_map)

    if not rows and builtin_profile != "Générique":
        with pdfplumber.open(BytesIO(content)) as pdf:
            page_texts = [p.extract_text() or "" for p in pdf.pages]
            full_text = "\n".join(page_texts)
            aliases_map = _aliases_from_rules(learned_rules, profile_key)
            rows = _extract_generic_pdf(pdf, owner, filename, full_text, aliases_map)
            if rows:
                profile += " / secours générique"

    return {
        "filename": filename,
        "profile": profile,
        "profile_key": profile_key,
        "profile_confidence": profile_confidence,
        "profile_source": profile_source,
        "signature": signature,
        "owner": owner,
        "row_count": len(rows),
        "rows": rows,
    }


def extract_customs_xlsx(filename: str, content: bytes, learned_profiles: List[Dict] | None = None, learned_rules: List[Dict] | None = None) -> Dict:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    rows_out = []
    signature = _signature_terms_xlsx(wb)
    matched_profile, learned_score = _best_learned_profile(signature, learned_profiles)
    if matched_profile:
        profile = _text(matched_profile.get("name")) or "Excel - profil appris"
        profile_key = _text(matched_profile.get("profile_key")) or _profile_key("douanes", profile)
        profile_confidence = learned_score
        profile_source = "learned"
    else:
        profile = "Excel - en-têtes reconnus"
        profile_key = f"douanes:excel:{_signature_hash(signature)}"
        profile_confidence = learned_score
        profile_source = "generic"
    aliases_map = _aliases_from_rules(learned_rules, profile_key)

    try:
        for ws in wb.worksheets:
            header_row = None
            mapping = {}
            header_values = {}
            for row_no in range(1, min(ws.max_row, 25) + 1):
                candidate = {}
                candidate_values = {}
                for col_no in range(1, ws.max_column + 1):
                    v = ws.cell(row_no, col_no).value
                    field = _match_header(v, aliases_map)
                    if field and field not in candidate:
                        candidate[field] = col_no
                        candidate_values[field] = _text(v)
                if "title" in candidate and len(candidate) >= 2:
                    header_row = row_no
                    mapping = candidate
                    header_values = candidate_values
                    break
            if not header_row:
                continue
            currency = _currency_from_header(header_values.get("value_source", ""), "")
            for row_no in range(header_row + 1, ws.max_row + 1):
                def raw_cell(field):
                    col = mapping.get(field)
                    return ws.cell(row_no, col).value if col else ""
                title = _text(raw_cell("title"))
                if not title:
                    continue
                ref = _text(raw_cell("reference"))
                raw = [None] * (max(mapping.values()) + 1 if mapping else 0)
                for field, col in mapping.items():
                    if col >= len(raw):
                        raw.extend([None] * (col - len(raw) + 1))
                    raw[col] = ws.cell(row_no, col).value
                zero_mapping = {field: col for field, col in mapping.items()}
                row = _base_row("", filename, row_no, ref)
                # openpyxl est indexé à partir de 1, notre tableau artificiel aussi.
                rows_out.append(_apply_generic_fields(row, zero_mapping, raw, header_values, currency))
    finally:
        wb.close()

    if not rows_out and not matched_profile:
        profile = "Excel - aucun tableau reconnu"

    return {
        "filename": filename,
        "profile": profile,
        "profile_key": profile_key,
        "profile_confidence": profile_confidence,
        "profile_source": profile_source,
        "signature": signature,
        "owner": "",
        "row_count": len(rows_out),
        "rows": rows_out,
    }


def extract_customs_document(filename: str, content: bytes, learned_profiles: List[Dict] | None = None, learned_rules: List[Dict] | None = None) -> Dict:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_customs_pdf(filename, content, learned_profiles, learned_rules)
    if lower.endswith(".xlsx"):
        return extract_customs_xlsx(filename, content, learned_profiles, learned_rules)
    raise ValueError("Format non encore pris en charge par le moteur DOUANES (PDF et XLSX disponibles dans cette version).")

# -----------------------------------------------------------------------------
# CLEANER ARTICLES - extraction Excel generique et profils appris
# -----------------------------------------------------------------------------
_ARTICLE_HEADER_ALIASES = {
    "reference": [
        "ref esi", "reference esi", "reference", "ref", "ref article",
        "inventory number", "inventory no", "object id", "item id", "accession number",
        "numero inventaire", "n inventaire",
    ],
    "quantite": ["quantite", "qte", "qty", "quantity", "nombre", "nb"],
    "longueur_cm": [
        "longueur cm", "longueur", "length cm", "lenght cm", "length", "lenght",
    ],
    "largeur_cm": ["largeur cm", "largeur", "width cm", "width"],
    "hauteur_cm": ["hauteur cm", "hauteur", "height cm", "height"],
    "poids_kg": [
        "poids brut kg", "poids kg", "poids", "gross weight kg", "gross weight",
        "weight kg", "weight", "kg",
    ],
}

_ARTICLE_DESIGNATION_TITLE_ALIASES = [
    "title", "artwork title", "object title", "object name", "designation", "product title",
]
_ARTICLE_DESIGNATION_DESCRIPTION_ALIASES = [
    "description", "product description", "description du produit", "object description",
]
_ARTICLE_TECH_ALIASES = {
    "artist": ["artist", "artiste", "maker", "creator"],
    "packing": ["packing", "emballage", "conditionnement"],
    "numero_caisse": ["numero de caisse", "numero caisse", "n caisse", "crate number", "crate no", "case number", "case no"],
    "medium_explicit": ["medium", "materials", "material", "technique"],
    "circa": ["circa", "date", "year", "annee"],
    "value_usd_source": ["value usd", "usd value", "value usd ", "customs value usd"],
    "value_eur_source": ["value eur", "value euro", "value €", "valeur eur", "valeur euro", "valeur €"],
}


def _article_aliases_from_rules(learned_rules: List[Dict] | None, profile_key: str = "") -> Dict[str, List[str]]:
    aliases = {k: list(v) for k, v in _ARTICLE_HEADER_ALIASES.items()}
    # Les regles apprises peuvent enrichir les champs de base ARTICLES.
    for rule in learned_rules or []:
        if not isinstance(rule, dict) or rule.get("active") is False:
            continue
        rule_profile = _text(rule.get("profile_key"))
        if rule_profile and profile_key and rule_profile != profile_key:
            continue
        if rule_profile and not profile_key:
            continue
        target = _text(rule.get("target_field"))
        alias = _text(rule.get("source_alias"))
        if target and alias:
            aliases.setdefault(target, [])
            if _norm(alias) not in {_norm(x) for x in aliases[target]}:
                aliases[target].append(alias)
    return aliases


def _article_header_match(value, aliases_map: Dict[str, List[str]]):
    """Retourne (champ, score) pour choisir le meilleur en-tete quand plusieurs sont possibles."""
    n = _norm(value)
    if not n:
        return "", 0
    best_field = ""
    best_score = 0
    for field, aliases in aliases_map.items():
        for pos, alias in enumerate(aliases):
            a = _norm(alias)
            if not a:
                continue
            if n == a:
                score = 100 - min(pos, 30)
            elif len(a) >= 5 and a in n:
                score = 60 - min(pos, 30)
            else:
                continue
            if score > best_score:
                best_field, best_score = field, score
    return best_field, best_score


def _article_aux_header_match(value):
    n = _norm(value)
    if not n:
        return "", 0

    # Title est prioritaire sur Description pour la designation quand les deux existent.
    for pos, alias in enumerate(_ARTICLE_DESIGNATION_TITLE_ALIASES):
        a = _norm(alias)
        if n == a or (len(a) >= 5 and a in n):
            return "title_source", 120 - min(pos, 20)
    for pos, alias in enumerate(_ARTICLE_DESIGNATION_DESCRIPTION_ALIASES):
        a = _norm(alias)
        if n == a or (len(a) >= 5 and a in n):
            return "description_source", 100 - min(pos, 20)

    for field, aliases in _ARTICLE_TECH_ALIASES.items():
        for pos, alias in enumerate(aliases):
            a = _norm(alias)
            if n == a or (len(a) >= 5 and a in n):
                return field, 90 - min(pos, 20)
    return "", 0


def _merged_value_lookup(ws):
    """Construit un petit index des cellules fusionnees -> valeur de la cellule maitre."""
    lookup = {}
    try:
        ranges = list(ws.merged_cells.ranges)
    except Exception:
        ranges = []
    for merged in ranges:
        value = ws.cell(merged.min_row, merged.min_col).value
        for row_no in range(merged.min_row, merged.max_row + 1):
            for col_no in range(merged.min_col, merged.max_col + 1):
                lookup[(row_no, col_no)] = value
    return lookup


def _ws_value(ws, row_no: int, col_no: int, merged_lookup: Dict):
    value = ws.cell(row_no, col_no).value
    if value is None and (row_no, col_no) in merged_lookup:
        return merged_lookup[(row_no, col_no)]
    return value


def _article_signature_xlsx(wb) -> List[str]:
    """Signature stable : noms d'onglets + libelles/metadata, jamais la liste complete des oeuvres."""
    terms = []
    keywords = (
        "packing list", "customer", "client", "file", "dossier", "ref", "reference",
        "title", "description", "designation", "qty", "quantity", "quantite",
        "length", "lenght", "longueur", "width", "largeur", "height", "hauteur",
        "weight", "poids", "kg", "artist", "packing", "caisse", "crate", "value",
    )
    for ws in wb.worksheets[:8]:
        if _norm(ws.title):
            terms.append("sheet " + _norm(ws.title))
        max_row = min(ws.max_row, 20)
        max_col = min(ws.max_column, 50)
        for row_no in range(1, max_row + 1):
            for col_no in range(1, max_col + 1):
                n = _norm(ws.cell(row_no, col_no).value)
                if not n or len(n) > 100:
                    continue
                if any(k in n for k in keywords):
                    terms.append(n)
    return sorted(set(terms))[:160]


def _article_builtin_profile(signature: List[str]) -> str:
    joined = " | ".join(signature)
    if "packing list" in joined and "pozzi" in joined and (
        "lenght cm" in joined or "ref esi" in joined or "customer pozzi" in joined
    ):
        return "Packing List Pozzi"
    return ""


def _detect_article_header_row(ws, aliases_map: Dict[str, List[str]], merged_lookup: Dict):
    """Cherche la meilleure ligne d'en-tetes dans les 40 premieres lignes de l'onglet."""
    best = None
    for row_no in range(1, min(ws.max_row, 40) + 1):
        mapping = {}
        mapping_scores = {}
        header_values = {}
        aux_mapping = {}
        aux_scores = {}
        aux_headers = {}

        for col_no in range(1, min(ws.max_column, 80) + 1):
            raw = _ws_value(ws, row_no, col_no, merged_lookup)
            txt = _text(raw)
            if not txt:
                continue

            field, score = _article_header_match(txt, aliases_map)
            if field and score > mapping_scores.get(field, -1):
                mapping[field] = col_no
                mapping_scores[field] = score
                header_values[field] = txt

            aux_field, aux_score = _article_aux_header_match(txt)
            if aux_field and aux_score > aux_scores.get(aux_field, -1):
                aux_mapping[aux_field] = col_no
                aux_scores[aux_field] = aux_score
                aux_headers[aux_field] = txt

        # Designation = Title en priorite, sinon Description.
        if "title_source" in aux_mapping:
            mapping["designation"] = aux_mapping["title_source"]
            mapping_scores["designation"] = aux_scores["title_source"]
            header_values["designation"] = aux_headers["title_source"]
        elif "description_source" in aux_mapping:
            mapping["designation"] = aux_mapping["description_source"]
            mapping_scores["designation"] = aux_scores["description_source"]
            header_values["designation"] = aux_headers["description_source"]

        # Une ligne d'en-tetes credible doit avoir au moins 3 champs et une ancre article.
        anchors = int("reference" in mapping) + int("designation" in mapping)
        if anchors == 0 or len(mapping) < 3:
            continue
        score = len(mapping) * 10 + anchors * 8 + sum(mapping_scores.values()) / 1000.0
        if best is None or score > best[0]:
            best = (score, row_no, mapping, header_values, aux_mapping, aux_headers)

    if not best:
        return None
    _, row_no, mapping, header_values, aux_mapping, aux_headers = best
    return {
        "row": row_no,
        "mapping": mapping,
        "header_values": header_values,
        "aux_mapping": aux_mapping,
        "aux_headers": aux_headers,
    }


def _article_number_text(value, default="") -> str:
    if value is None:
        return default
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return (f"{value:.10f}").rstrip("0").rstrip(".")
    return _text(value)


def extract_articles_xlsx(
    filename: str,
    content: bytes,
    learned_profiles: List[Dict] | None = None,
    learned_rules: List[Dict] | None = None,
) -> Dict:
    """Extrait une liste d'articles d'un classeur Excel heterogene, sans imposer la ligne 1."""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(content), read_only=False, data_only=True)
    rows_out = []
    sheets_used = []
    detected_headers = []
    signature = _article_signature_xlsx(wb)

    builtin_profile = _article_builtin_profile(signature)
    matched_profile = None
    learned_score = 0.0
    if not builtin_profile:
        matched_profile, learned_score = _best_learned_profile(signature, learned_profiles)

    if builtin_profile:
        profile = builtin_profile
        profile_key = _profile_key("articles", profile)
        profile_confidence = 1.0
        profile_source = "built-in"
    elif matched_profile:
        profile = _text(matched_profile.get("name")) or "Excel - profil ARTICLES appris"
        profile_key = _text(matched_profile.get("profile_key")) or _profile_key("articles", profile)
        profile_confidence = learned_score
        profile_source = "learned"
    else:
        profile = "Excel ARTICLES - structure reconnue"
        profile_key = f"articles:excel:{_signature_hash(signature)}"
        profile_confidence = learned_score
        profile_source = "generic"

    aliases_map = _article_aliases_from_rules(learned_rules, profile_key)

    try:
        for ws in wb.worksheets:
            merged_lookup = _merged_value_lookup(ws)
            detected = _detect_article_header_row(ws, aliases_map, merged_lookup)
            if not detected:
                continue

            header_row = detected["row"]
            mapping = detected["mapping"]
            header_values = detected["header_values"]
            aux_mapping = detected["aux_mapping"]
            aux_headers = detected["aux_headers"]
            detected_headers.append({
                "sheet": ws.title,
                "row": header_row,
                "headers": dict(header_values),
            })

            sheet_count_before = len(rows_out)
            empty_streak = 0
            for row_no in range(header_row + 1, ws.max_row + 1):
                def cell(field):
                    col = mapping.get(field)
                    return _ws_value(ws, row_no, col, merged_lookup) if col else None

                def aux_cell(field):
                    col = aux_mapping.get(field)
                    return _ws_value(ws, row_no, col, merged_lookup) if col else None

                reference = _article_number_text(cell("reference"))
                designation = _text(cell("designation"))

                # Ignore les lignes TOTAL / sous-totaux et les zones hors liste.
                if _norm(reference) in {"total", "totaux"} or _norm(designation) in {"total", "totaux"}:
                    continue
                if not reference and not designation:
                    empty_streak += 1
                    # On continue pour tolerer des espaces dans les feuilles, sans parcourir 1000 lignes vides.
                    if empty_streak >= 20:
                        break
                    continue
                empty_streak = 0

                qty = _article_number_text(cell("quantite"), "1") or "1"
                longueur = _article_number_text(cell("longueur_cm"))
                largeur = _article_number_text(cell("largeur_cm"))
                hauteur = _article_number_text(cell("hauteur_cm"))
                poids = _article_number_text(cell("poids_kg"))

                # Donnees techniques utiles pour de futurs champs personnalises.
                title_source = _text(aux_cell("title_source"))
                description_source = _text(aux_cell("description_source"))
                explicit_medium = _text(aux_cell("medium_explicit"))
                medium = explicit_medium or (description_source if title_source else "")
                artist = _text(aux_cell("artist"))
                packing = _text(aux_cell("packing"))
                numero_caisse = _text(aux_cell("numero_caisse"))
                circa = _text(aux_cell("circa"))

                value_usd_raw = aux_cell("value_usd_source")
                value_eur_raw = aux_cell("value_eur_source")
                source_value = ""
                source_value_number = None
                source_currency = ""
                if value_usd_raw not in (None, ""):
                    source_value, source_value_number = _money(value_usd_raw)
                    source_currency = "USD"
                elif value_eur_raw not in (None, ""):
                    source_value, source_value_number = _money(value_eur_raw)
                    source_currency = "EUR"

                row = {
                    "reference": reference,
                    "designation": designation,
                    "quantite": qty,
                    "longueur_cm": longueur,
                    "largeur_cm": largeur,
                    "hauteur_cm": hauteur,
                    "poids_kg": poids,
                    "artist": artist,
                    "medium": medium,
                    "packing": packing,
                    "numero_caisse": numero_caisse,
                    "circa": circa,
                    "source_value": source_value,
                    "source_value_number": source_value_number,
                    "source_currency": source_currency,
                    "source_file": filename,
                    "source_sheet": ws.title,
                    "source_page": row_no,
                    "source_headers": dict(header_values),
                    "_profile_key": profile_key,
                }

                # La designation peut venir de Title ou Description selon l'onglet.
                if "designation" in header_values:
                    row["source_headers"]["designation"] = header_values["designation"]
                rows_out.append(row)

            if len(rows_out) > sheet_count_before:
                sheets_used.append({
                    "name": ws.title,
                    "header_row": header_row,
                    "row_count": len(rows_out) - sheet_count_before,
                })
    finally:
        wb.close()

    if not rows_out and not matched_profile and not builtin_profile:
        profile = "Excel ARTICLES - aucun tableau reconnu"

    return {
        "filename": filename,
        "profile": profile,
        "profile_key": profile_key,
        "profile_confidence": profile_confidence,
        "profile_source": profile_source,
        "signature": signature,
        "owner": "",
        "row_count": len(rows_out),
        "rows": rows_out,
        "sheets": sheets_used,
        "detected_headers": detected_headers,
    }


def _split_dimensions_cm(value) -> Tuple[str, str, str]:
    """Extrait L x l x H depuis une dimension libre, sans inventer de valeur."""
    raw = _text(value)
    if not raw:
        return "", "", ""
    nums = re.findall(r"[-+]?\d+(?:[.,]\d+)?", raw)
    if len(nums) < 3:
        return "", "", ""
    vals = [x.replace(",", ".") for x in nums[:3]]
    return vals[0], vals[1], vals[2]


def _article_pdf_signature(pdf, full_text: str) -> List[str]:
    terms = []
    for raw in (full_text or "").replace("\r", "").splitlines()[:180]:
        n = _norm(raw)
        if not n or len(n) > 120:
            continue
        if any(k in n for k in (
            "packing list", "liste de colisage", "colis", "reference", "description",
            "quantite", "dimensions", "poids net", "poids brut", "gross weight",
            "net weight", "jouffre", "incoterm", "mode de transport",
        )):
            terms.append(n)
    for page in list(pdf.pages)[:4]:
        try:
            tables = page.extract_tables() or []
        except Exception:
            tables = []
        for table in tables[:3]:
            for row in (table or [])[:12]:
                for cell in row or []:
                    n = _norm(cell)
                    if n and 2 <= len(n) <= 100 and any(k in n for k in (
                        "colis", "reference", "description", "quantite", "dimensions", "poids", "weight"
                    )):
                        terms.append(n)
    return sorted(set(terms))[:160]


def _extract_jouffre_articles_pdf(pdf, filename: str, profile_key: str) -> List[Dict]:
    """Profil robuste pour les listes de colisage Jouffre type VCO."""
    rows_out = []
    current_colis = ""
    current_zone = ""

    for page_no, page in enumerate(pdf.pages, 1):
        try:
            tables = page.extract_tables() or []
        except Exception:
            tables = []
        for table in tables:
            for raw in table or []:
                raw = list(raw or []) + [None] * 10
                first = _text(raw[0])
                first_norm = _norm(first)
                if not first and not any(_text(x) for x in raw[1:9]):
                    continue
                if "totaux" in first_norm or "nombre total de colis" in first_norm:
                    continue
                if first_norm in {"colis n", "colis no", "colis"}:
                    continue

                if re.fullmatch(r"\d+", first):
                    current_colis = first
                elif first and not re.fullmatch(r"\d+", first):
                    # En-tetes et blocs administratifs de la premiere page.
                    continue

                zone = _text(raw[1])
                if zone:
                    current_zone = zone
                description = _text(raw[2])
                reference = _text(raw[3])
                qty_main = _article_number_text(raw[4])
                qty_cushions = _article_number_text(raw[5])
                dims_raw = _text(raw[6])
                net_weight = _article_number_text(raw[7])
                gross_weight = _article_number_text(raw[8])

                # Une ligne article doit au minimum contenir une description ou une reference.
                if not description and not reference:
                    continue
                if not current_colis:
                    continue

                longueur, largeur, hauteur = _split_dimensions_cm(dims_raw)
                qty = qty_main or "1"
                row = {
                    "reference": reference,
                    "designation": description,
                    "quantite": qty,
                    "longueur_cm": longueur,
                    "largeur_cm": largeur,
                    "hauteur_cm": hauteur,
                    "poids_kg": net_weight,
                    "poids_brut_kg": gross_weight,
                    "quantite_coussins": qty_cushions,
                    "colis": current_colis,
                    "zone": current_zone,
                    "dimensions": _dimensions(dims_raw, "cm") if dims_raw else "",
                    "observation": current_zone,
                    "source_file": filename,
                    "source_page": page_no,
                    "source_headers": {
                        "reference": "Référence",
                        "designation": "Description",
                        "quantite": "Quantité Canapés/Fauteuils/Voilages/Rideaux",
                        "dimensions": "Dimensions (Cm)",
                        "poids_kg": "Poids net (kg)",
                    },
                    "_profile_key": profile_key,
                }
                rows_out.append(row)
    return rows_out


def _extract_generic_articles_pdf(pdf, filename: str, profile_key: str) -> List[Dict]:
    """Secours PDF generique pour tableaux simples ARTICLES/PACKING LIST."""
    rows_out = []
    current_colis = ""
    for page_no, page in enumerate(pdf.pages, 1):
        try:
            tables = page.extract_tables() or []
        except Exception:
            tables = []
        for table in tables:
            if not table:
                continue
            best = None
            for i, raw in enumerate((table or [])[:10]):
                mapping = {}
                headers = {}
                for col, value in enumerate(raw or []):
                    n = _norm(value)
                    if not n:
                        continue
                    if n in {"reference", "ref", "ref esi", "item id", "object id"} or "reference" in n:
                        mapping.setdefault("reference", col); headers.setdefault("reference", _text(value))
                    elif n in {"description", "designation", "title", "object name"} or "description" in n:
                        mapping.setdefault("designation", col); headers.setdefault("designation", _text(value))
                    elif n in {"quantite", "qte", "qty", "quantity"} or n.startswith("quantite "):
                        mapping.setdefault("quantite", col); headers.setdefault("quantite", _text(value))
                    elif "dimension" in n or n in {"dims", "size"}:
                        mapping.setdefault("dimensions", col); headers.setdefault("dimensions", _text(value))
                    elif "poids net" in n or "net weight" in n:
                        mapping.setdefault("poids_kg", col); headers.setdefault("poids_kg", _text(value))
                    elif "poids brut" in n or "gross weight" in n:
                        mapping.setdefault("poids_brut_kg", col); headers.setdefault("poids_brut_kg", _text(value))
                    elif n in {"colis", "colis n", "colis no", "crate", "case", "package"} or n.startswith("colis "):
                        mapping.setdefault("colis", col); headers.setdefault("colis", _text(value))
                score = len(mapping) + (2 if "designation" in mapping else 0) + (1 if "reference" in mapping else 0)
                if score >= 4 and (best is None or score > best[0]):
                    best = (score, i, mapping, headers)
            if not best:
                continue
            _, header_idx, mapping, headers = best
            for raw in table[header_idx + 1:]:
                raw = list(raw or [])
                def cell(name):
                    idx = mapping.get(name)
                    return raw[idx] if idx is not None and idx < len(raw) else None
                colis = _text(cell("colis"))
                if colis:
                    if "total" in _norm(colis):
                        continue
                    current_colis = colis
                designation = _text(cell("designation"))
                reference = _text(cell("reference"))
                if not designation and not reference:
                    continue
                dims_raw = _text(cell("dimensions"))
                longueur, largeur, hauteur = _split_dimensions_cm(dims_raw)
                row = {
                    "reference": reference,
                    "designation": designation,
                    "quantite": _article_number_text(cell("quantite"), "1") or "1",
                    "longueur_cm": longueur,
                    "largeur_cm": largeur,
                    "hauteur_cm": hauteur,
                    "poids_kg": _article_number_text(cell("poids_kg")),
                    "poids_brut_kg": _article_number_text(cell("poids_brut_kg")),
                    "colis": current_colis,
                    "dimensions": _dimensions(dims_raw, "cm") if dims_raw else "",
                    "observation": "",
                    "source_file": filename,
                    "source_page": page_no,
                    "source_headers": headers,
                    "_profile_key": profile_key,
                }
                rows_out.append(row)
    return rows_out


def extract_articles_pdf(
    filename: str,
    content: bytes,
    learned_profiles: List[Dict] | None = None,
    learned_rules: List[Dict] | None = None,
) -> Dict:
    import pdfplumber

    with pdfplumber.open(BytesIO(content)) as pdf:
        full_text = "\n".join((page.extract_text() or "") for page in pdf.pages)
        signature = _article_pdf_signature(pdf, full_text)
        n = _norm(full_text)
        builtin_profile = ""
        if "jouffre maroc" in n and "liste de colisage" in n:
            builtin_profile = "Jouffre - Liste de colisage VCO"

        matched_profile = None
        learned_score = 0.0
        if not builtin_profile:
            matched_profile, learned_score = _best_learned_profile(signature, learned_profiles)

        if builtin_profile:
            profile = builtin_profile
            profile_key = _profile_key("articles", profile)
            profile_confidence = 1.0
            profile_source = "built-in"
        elif matched_profile:
            profile = _text(matched_profile.get("name")) or "PDF - profil ARTICLES appris"
            profile_key = _text(matched_profile.get("profile_key")) or _profile_key("articles", profile)
            profile_confidence = learned_score
            profile_source = "learned"
        else:
            profile = "PDF ARTICLES - structure reconnue"
            profile_key = f"articles:pdf:{_signature_hash(signature)}"
            profile_confidence = learned_score
            profile_source = "generic"

        if builtin_profile:
            rows = _extract_jouffre_articles_pdf(pdf, filename, profile_key)
        else:
            rows = _extract_generic_articles_pdf(pdf, filename, profile_key)

    if not rows and not builtin_profile and not matched_profile:
        profile = "PDF ARTICLES - aucun tableau reconnu"

    return {
        "filename": filename,
        "profile": profile,
        "profile_key": profile_key,
        "profile_confidence": profile_confidence,
        "profile_source": profile_source,
        "signature": signature,
        "owner": "",
        "row_count": len(rows),
        "rows": rows,
        "sheets": [],
        "detected_headers": [],
    }


def extract_articles_document(
    filename: str,
    content: bytes,
    learned_profiles: List[Dict] | None = None,
    learned_rules: List[Dict] | None = None,
) -> Dict:
    lower = filename.lower()
    if lower.endswith(".xlsx"):
        return extract_articles_xlsx(filename, content, learned_profiles, learned_rules)
    if lower.endswith(".pdf"):
        return extract_articles_pdf(filename, content, learned_profiles, learned_rules)
    raise ValueError("Format non encore pris en charge par le moteur ARTICLES (PDF et XLSX disponibles dans cette version).")
